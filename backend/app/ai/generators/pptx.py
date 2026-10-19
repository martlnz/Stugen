"""Xuất SlideDeck ra PPTX bằng python-pptx (ADR-0011); LaTeX giữ dạng văn bản thô."""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Inches, Pt

from app.schemas.llm import Slide, SlideDeck

# Layout của template mặc định python-pptx
LAYOUT_TITLE, LAYOUT_CONTENT, LAYOUT_SECTION = 0, 1, 2
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
EXAMPLE_FILL = RGBColor(0xEE, 0xF3, 0xFB)


def _add_notes(pptx_slide, slide: Slide) -> None:  # type: ignore[no-untyped-def]
    notes = slide.notes or ""
    if slide.page_refs:
        refs = ", ".join(str(p) for p in slide.page_refs)
        notes = f"{notes}\n\nTrang tham chiếu: {refs}".strip()
    if notes:
        pptx_slide.notes_slide.notes_text_frame.text = notes


def _content_slide(prs, slide: Slide) -> object:  # type: ignore[no-untyped-def]
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    s.shapes.title.text = slide.title
    width, height = prs.slide_width, prs.slide_height
    body = s.placeholders[1]
    body.left, body.top = Inches(0.6), Inches(1.5)
    body.width = width - Inches(1.2)
    example_h = Inches(1.5) if slide.example else 0
    body.height = height - body.top - Inches(0.6) - example_h
    frame = body.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    size = Pt(20 if len(slide.bullets) <= 4 else 18)
    for i, bullet in enumerate(slide.bullets):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = bullet
        para.font.size = size
    if slide.explanation:
        para = frame.add_paragraph()
        para.text = slide.explanation
        para.font.size = Pt(16)
        para.font.italic = True
        para.font.color.rgb = MUTED
        para.space_before = Pt(10)
    if slide.example:
        box = s.shapes.add_textbox(
            Inches(0.6), height - Inches(0.5) - example_h, width - Inches(1.2), example_h
        )
        box.fill.solid()
        box.fill.fore_color.rgb = EXAMPLE_FILL
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        head = tf.paragraphs[0].add_run()
        head.text = "Ví dụ: "
        head.font.bold = True
        head.font.size = Pt(15)
        head.font.color.rgb = ACCENT
        run = tf.paragraphs[0].add_run()
        run.text = slide.example.removeprefix("Ví dụ minh hoạ:").strip()
        run.font.size = Pt(15)
    if slide.page_refs:
        box = s.shapes.add_textbox(Inches(0.6), height - Inches(0.45), width, Pt(18))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Trang " + ", ".join(str(p) for p in slide.page_refs)
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED
    return s


def build_pptx(deck: SlideDeck) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    width = prs.slide_width

    for slide in deck.slides:
        if slide.layout in ("title", "section"):
            layout = LAYOUT_TITLE if slide.layout == "title" else LAYOUT_SECTION
            s = prs.slides.add_slide(prs.slide_layouts[layout])
            s.shapes.title.text = slide.title
            s.placeholders[1].text = slide.subtitle or ""
        else:
            s = _content_slide(prs, slide)
        title = s.shapes.title  # type: ignore[attr-defined]
        title.text_frame.paragraphs[0].font.color.rgb = ACCENT
        title.width = width - title.left * 2
        _add_notes(s, slide)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
